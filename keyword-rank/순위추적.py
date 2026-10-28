# -*- coding: utf-8 -*-
"""센터 키워드 순위 추적기 (센터 이름은 설정_블로그.json 「센터_표시명」에서 읽는다)

매주 실행하면:
1. 네이버 검색광고 API로 키워드별 검색량(PC/모바일) 수집
2. 네이버 블로그 탭 상위 30개에서 선생님 글 순위 + 리뷰어(외부) 글 순위 판정
3. 네이버 플레이스에서 센터 순위(광고 제외) 수집
4. 키워드순위표.xlsx 에 날짜 시트로 저장 (직전 기록과 변동 표시)

사용법:  python 순위추적.py            # 전체 실행
        python 순위추적.py 원주pt      # 특정 키워드만 (테스트용)
"""
import base64
import hashlib
import hmac
import json
import re
import sys
import time
import urllib.parse
import urllib.request
from datetime import date, datetime
from email.utils import parsedate_to_datetime
from pathlib import Path

HERE = Path(__file__).parent
API_CONFIG = HERE / "_네이버광고_API.json"
KEYWORDS_FILE = HERE / "설정_키워드.txt"
BLOGS_FILE = HERE / "설정_블로그.json"
HISTORY_DIR = HERE / "기록"
XLSX_FILE = HERE / "키워드순위표.xlsx"

def 센터표시():
    """표 머리·안내에 찍을 센터 이름.

    2026-10-27 — 다른 센터 대표에게도 주게 되어, 코드에 박혀 있던 「비브레PT혁신점」을 설정으로 뺐다.
    설정_블로그.json 「센터_표시명」 → 없으면 「플레이스_센터명 + 지점」(옛 설치본은 이 칸이 없다)."""
    try:
        with open(BLOGS_FILE, encoding="utf-8") as f:
            c = json.load(f)
    except Exception:
        return "우리 센터"
    if (c.get("센터_표시명") or "").strip():
        return c["센터_표시명"].strip()
    이름, 지점 = (c.get("플레이스_센터명") or "").strip(), (c.get("플레이스_지점_우선") or "").strip()
    return f"{이름} {지점}점" if 이름 and 지점 else (이름 or "우리 센터")


BLOG_TOP_N = 30          # 블로그 탭 수집 상한
PLACE_TOP_N = 70         # 플레이스 1쪽(광고 제외 70곳)

# 순위별 클릭률 추정 (노출점수 계산용, 대략치)
CTR = {1: .25, 2: .13, 3: .09, 4: .06, 5: .05, 6: .04, 7: .035, 8: .03, 9: .025, 10: .02}


def ctr_of(rank):
    if not rank:
        return 0.0
    if rank in CTR:
        return CTR[rank]
    return .01 if rank <= 20 else .005


def norm(s):
    return re.sub(r"\s+", "", s).upper()


# ──────────────────────────── 1) 검색량 (검색광고 API) ────────────────────────────

def fetch_volumes(keywords, cfg):
    base = "https://api.searchad.naver.com"
    uri = "/keywordstool"

    # PC 시계가 어긋나 있어도 되도록 서버 Date 헤더로 보정
    probe = urllib.request.Request(base, method="HEAD")
    try:
        with urllib.request.urlopen(probe, timeout=10) as r:
            server_dt = parsedate_to_datetime(r.headers["Date"])
    except urllib.error.HTTPError as e:
        server_dt = parsedate_to_datetime(e.headers["Date"])
    offset_ms = round(server_dt.timestamp() * 1000) - round(time.time() * 1000)
    if abs(offset_ms) > 3000:
        print(f"  [시계 보정] PC가 서버와 {offset_ms/1000:+.1f}초 어긋남 → 보정")

    def parse_cnt(v):
        # "< 10" 은 0으로 (대표님 표 관례)
        if isinstance(v, str) and "<" in v:
            return 0
        try:
            return int(v)
        except (TypeError, ValueError):
            return 0

    volumes = {}
    related = {}  # 연관 키워드 풀 (신규 키워드 제안용)
    for i in range(0, len(keywords), 5):
        batch = [k.replace(" ", "") for k in keywords[i:i + 5]]
        timestamp = str(round(time.time() * 1000) + offset_ms)
        msg = f"{timestamp}.GET.{uri}"
        sig = base64.b64encode(
            hmac.new(cfg["secret_key"].encode(), msg.encode(), hashlib.sha256).digest()
        ).decode()
        params = urllib.parse.urlencode({"hintKeywords": ",".join(batch), "showDetail": "1"})
        req = urllib.request.Request(f"{base}{uri}?{params}", headers={
            "X-Timestamp": timestamp,
            "X-API-KEY": cfg["access_license"],
            "X-Customer": str(cfg["customer_id"]),
            "X-Signature": sig,
        })
        with urllib.request.urlopen(req, timeout=15) as resp:
            data = json.loads(resp.read().decode())
        want = {norm(k) for k in batch}
        for row in data.get("keywordList", []):
            key = norm(row["relKeyword"])
            pcv = parse_cnt(row["monthlyPcQcCnt"])
            mov = parse_cnt(row["monthlyMobileQcCnt"])
            if key in want:
                if key not in volumes:
                    volumes[key] = (pcv, mov)
            elif key not in related:
                related[key] = {"키워드": row["relKeyword"], "PC": pcv, "모바일": mov,
                                "합계": pcv + mov, "광고경쟁": row.get("compIdx", "")}
        time.sleep(0.3)
    return volumes, related


# ──────────────────────────── 2) 블로그 탭 + 플레이스 (Playwright) ────────────────────────────

EXTRACT_POSTS_JS = """() => {
    const root = document.querySelector('ul.lst_view') || document.body;
    const seen = new Set();
    const posts = [];
    for (const a of root.querySelectorAll('a[href*="blog.naver.com"]')) {
        const m = a.href.match(/blog\\.naver\\.com\\/([A-Za-z0-9_-]+)\\/(\\d{8,})/);
        if (!m) continue;
        const key = m[1] + '/' + m[2];
        const text = (a.innerText || '').replace(/새 창 열림/g, '').replace(/\\s+/g, ' ').trim();
        if (!seen.has(key)) {
            seen.add(key);
            posts.push({blog: m[1], post: m[2], title: text, desc: ''});
        } else if (text.length > 5) {
            const p = posts.find(x => x.blog + '/' + x.post === key);
            if (p) { if (!p.title) p.title = text; else if (text !== p.title && text.length > p.desc.length) p.desc = text; }
        }
    }
    return posts;
}"""

EXTRACT_PLACES_JS = """() => {
    const s = window.__APOLLO_STATE__ || {};
    const rq = s.ROOT_QUERY || {};
    const listKey = Object.keys(rq).find(k => k.startsWith('placeList'));
    if (!listKey || !rq[listKey] || !rq[listKey].businesses) return [];
    const biz = rq[listKey].businesses;
    const items = Array.isArray(biz) ? biz : (biz.items || []);
    const out = [];
    for (const ref of items) {
        const item = ref && ref.__ref ? s[ref.__ref] : null;
        if (item && item.name) out.push({id: item.id || '', name: item.name});
    }
    return out;
}"""


NO_RESULT_JS = """() => /검색\\s*결과가\\s*없습니다|조건에 맞는 업체가 없습니다/.test(document.body.innerText || '')"""


def read_places(page, kw):
    """플레이스 순위 목록을 읽는다. (목록, 오류문구 또는 None)

    ⚠ 2026-10-27 대표님 지적으로 고침 — 「원주pt 4위인데 원주피티는 이탈」
       원주pt·원주피티는 네이버가 같은 검색으로 보아 순위가 늘 같다. 그런데 표에선 한쪽만 비었다.
       원인: 지도를 연 뒤 **3.5초만** 기다리고 목록을 읽었다. 그 안에 지도가 다 안 뜨면
       프레임이 없거나 목록이 빈 채로 읽혔고, 코드는 그걸 「플레이스 결과가 원래 없는 키워드」로
       보고 **오류도 안 남겼다.** 그래서 순위가 비고, 지난주와 견주며 「이탈(전2)」로 찍혔다.
       캡차는 「확인못함」으로 갈라 뒀는데, 느리게 뜬 경우는 그 길을 안 탔다.
    고친 것
       · 목록이 실제로 뜰 때까지 최대 12초 지켜본다
       · 목록도 「검색 결과 없음」 문구도 안 보이면 한 번 더 연다
       · 그래도 못 받으면 오류를 남긴다 → 표에 「확인못함」(순위 하락 아님)
       · 「검색 결과가 없습니다」가 보일 때만 진짜 결과 없음으로 본다
    """
    url = f"https://map.naver.com/p/search/{urllib.parse.quote(kw)}"
    for attempt in range(2):
        page.goto(url, wait_until="domcontentloaded", timeout=30000)
        deadline = time.time() + 12
        while time.time() < deadline:
            page.wait_for_timeout(700)
            frame = next((f for f in page.frames if "pcmap.place.naver.com" in f.url), None)
            if not frame:
                continue
            try:
                if "보안문자" in frame.content():
                    return [], "플레이스 캡차"
                places = frame.evaluate(EXTRACT_PLACES_JS)
                if places:
                    return places[:PLACE_TOP_N], None
                if frame.evaluate(NO_RESULT_JS):
                    return [], None                      # 진짜로 플레이스 결과가 없는 키워드
            except Exception:
                continue                                 # 프레임이 갈아끼워지는 중 — 다시 본다
        if attempt == 0:
            page.wait_for_timeout(1500)
    return [], "플레이스 목록을 못 받음(지도 로딩 지연)"


def collect_serp(keywords, page):
    """키워드별 {blog_posts, places} 수집. page 는 Playwright page."""
    results = {}
    for idx, kw in enumerate(keywords, 1):
        entry = {"blog_posts": [], "places": [], "errors": []}

        # ── 블로그 탭 ──
        try:
            url = f"https://search.naver.com/search.naver?ssc=tab.blog.all&sm=tab_jum&query={urllib.parse.quote(kw)}"
            page.goto(url, wait_until="domcontentloaded", timeout=30000)
            page.wait_for_timeout(2000)
            if "보안문자" in page.content():
                entry["errors"].append("블로그탭 캡차")
            else:
                posts = page.evaluate(EXTRACT_POSTS_JS)
                stall = 0
                while len(posts) < BLOG_TOP_N and stall < 3:
                    page.mouse.wheel(0, 5000)
                    page.wait_for_timeout(1100)
                    new_posts = page.evaluate(EXTRACT_POSTS_JS)
                    stall = stall + 1 if len(new_posts) == len(posts) else 0
                    posts = new_posts
                entry["blog_posts"] = posts[:BLOG_TOP_N]
                # 블로그 탭이 한 건도 없으면 로딩 실패다 — 우리 키워드는 블로그 글이 늘 있다.
                # 빈 채로 두면 선생님 순위가 전부 「이탈」로 찍힌다 (플레이스와 같은 함정, 2026-10-27)
                if not posts:
                    entry["errors"].append("블로그탭 목록 0(로딩 실패)")
        except Exception as e:
            entry["errors"].append(f"블로그탭 {type(e).__name__}: {e}")

        # ── 플레이스 ──
        try:
            entry["places"], err = read_places(page, kw)
            if err:
                entry["errors"].append(err)
            # 그때 읽은 지도 주소(지도 중심·검색어가 들어 있다)를 남긴다 — 이상한 순위가 나오면
            # 「그때 무슨 목록을 받았나」를 나중에 확인할 수 있게 (2026-10-28, 원주헬스PT 1위 오판 조사)
            fr = next((f for f in page.frames if "pcmap.place.naver.com" in f.url), None)
            entry["place_frame"] = fr.url[:300] if fr else ""
        except Exception as e:
            entry["errors"].append(f"플레이스 {type(e).__name__}: {e}")

        results[kw] = entry
        print(f"  ({idx}/{len(keywords)}) {kw}: 블로그 {len(entry['blog_posts'])}건, "
              f"플레이스 {len(entry['places'])}곳"
              + (f"  ⚠ {entry['errors']}" if entry["errors"] else ""))
        time.sleep(0.8)

    # ── 마지막에 한 번 더 — 플레이스를 못 받은 키워드만 ──
    # 그 순간 네트워크가 느렸던 것이면 몇 분 뒤엔 대개 된다. 캡차는 다시 두드리지 않는다
    again = [kw for kw, e in results.items()
             if any(x.startswith("플레이스 목록을 못 받음") for x in e["errors"])]
    if again:
        print(f"  ↻ 플레이스를 못 받은 {len(again)}개 다시 시도: {', '.join(again)}")
        for kw in again:
            try:
                places, err = read_places(page, kw)
            except Exception as e:
                places, err = [], f"플레이스 {type(e).__name__}: {e}"
            if not err:
                results[kw]["places"] = places
                results[kw]["errors"] = [x for x in results[kw]["errors"]
                                         if not x.startswith("플레이스 목록을 못 받음")]
                print(f"    {kw}: 플레이스 {len(places)}곳 ✓")
            else:
                print(f"    {kw}: 또 실패 — 표에는 「확인못함」으로 남깁니다")
            time.sleep(0.8)
    return results


# ──────────────────────────── 3) 판정 ────────────────────────────

UA_MOBILE = ("Mozilla/5.0 (Linux; Android 13; SM-S908N) AppleWebKit/537.36 "
             "(KHTML, like Gecko) Chrome/124.0.0.0 Mobile Safari/537.36")
CACHE_FILE = HISTORY_DIR / "_본문판정캐시.json"


def check_post_content(post, blog_cfg, cache, stats):
    """제목·요약으로 판정 안 된 글의 본문을 직접 열어 센터 관련 여부 판정.
    반환: '지도'(플레이스 첨부 일치) | '본문'(언급) | '무관' | '확인실패'"""
    key = f"{post['blog']}/{post['post']}"
    if key in cache:
        return cache[key]["판정"]
    try:
        req = urllib.request.Request(f"https://m.blog.naver.com/{key}",
                                     headers={"User-Agent": UA_MOBILE})
        with urllib.request.urlopen(req, timeout=15) as r:
            html = r.read().decode("utf-8", errors="ignore")
    except Exception:
        stats["실패"] += 1
        return "확인실패"  # 실패는 캐시하지 않음 — 다음 실행 때 재시도
    place_id = blog_cfg.get("플레이스ID", "")
    if place_id and place_id in html:
        verdict = "지도"
    elif any(w.lower() in html.lower() for w in blog_cfg["센터_언급_단어"]):
        verdict = "본문"
    else:
        verdict = "무관"
    cache[key] = {"판정": verdict, "확인일": date.today().isoformat()}
    stats["조회"] += 1
    time.sleep(0.15)
    return verdict


def analyze(kw, entry, volumes, blog_cfg, cache, stats):
    trainers = blog_cfg["선생님"]
    mention_words = [w.lower() for w in blog_cfg["센터_언급_단어"]]

    pc, mo = volumes.get(norm(kw), (0, 0))
    row = {"키워드": kw, "PC": pc, "모바일": mo, "합계": pc + mo,
           "선생님": {}, "선생님상세": [], "리뷰어": [], "리뷰어글": [],
           "플레이스": None, "플레이스명": "", "오류": entry["errors"],
           # 플레이스 '측정 실패'(캡차·지도 로딩 오류)와 '진짜 순위권 밖'을 구분하는 재료.
           # 실패면 플레이스=None 이 "권외"가 아니므로, 다음 주 비교에서 신규/이탈로 단정하면 안 된다.
           "플레이스오류": any("플레이스" in e for e in entry["errors"]),
           # ⚠ 2026-10-28 대표님 지적 — 원주헬스PT 가 몇 주째 1위로 적혔는데 폰에선 10위였다.
           #   다시 재면 지도 위치·줌·키워드 순서를 어떻게 바꿔도 10위였고, 그때 받은 목록이
           #   안 남아 있어 원인을 못 가렸다. 그래서 받은 순위표 앞 15곳과 지도 주소를 같이 남긴다.
           #   다음에 이상한 순위가 보이면 기록/날짜.json 의 이 칸을 열어 보면 된다.
           "플레이스상위": [[i + 1, p.get("id", ""), p.get("name", "")]
                         for i, p in enumerate(entry.get("places", [])[:15])],
           "플레이스주소": entry.get("place_frame", "")}

    for rank, post in enumerate(entry["blog_posts"], 1):
        bid = post["blog"].lower()
        text = (post["title"] + " " + post["desc"]).lower()
        owner = next((name for name, ids in trainers.items() if bid in [i.lower() for i in ids]), None)
        if owner:
            row["선생님"].setdefault(owner, []).append(rank)
            row["선생님상세"].append({"선생님": owner, "블로그": bid, "글": post["post"],
                                    "순위": rank, "키워드": kw})
            continue
        if any(w in text for w in mention_words):
            basis = "제목"
        else:
            basis = check_post_content(post, blog_cfg, cache, stats)
            if basis in ("무관", "확인실패"):
                continue
        row["리뷰어"].append(rank)
        row["리뷰어글"].append({"순위": rank, "블로그": post["blog"],
                              "제목": post["title"][:60], "근거": basis})

    matches = [(i + 1, p["name"]) for i, p in enumerate(entry["places"])
               if blog_cfg["플레이스_센터명"] in p["name"]]
    if matches:
        preferred = [m for m in matches if blog_cfg["플레이스_지점_우선"] in m[1]]
        row["플레이스"], row["플레이스명"] = (preferred or matches)[0]

    best_blog = min((r for ranks in row["선생님"].values() for r in ranks), default=None)
    row["노출점수"] = round(row["합계"] * ctr_of(best_blog) + row["합계"] * ctr_of(row["플레이스"]) * 0.5)
    return row


# ──────────────────────────── 3.5) 선생님 블로그 지표 (전부 실측 기반) ────────────────────────────

def fetch_rss(blog_id):
    """네이버 공식 RSS에서 최근 글 목록. 반환: [{글번호, 발행일(datetime)}]"""
    import xml.etree.ElementTree as ET
    req = urllib.request.Request(f"https://rss.blog.naver.com/{blog_id}.xml",
                                 headers={"User-Agent": UA_MOBILE})
    with urllib.request.urlopen(req, timeout=15) as r:
        root = ET.fromstring(r.read())
    out = []
    for item in root.iter("item"):
        link = item.findtext("link") or ""
        m = re.search(r"/(\d{8,})", link)
        pub = item.findtext("pubDate")
        if m and pub:
            out.append({"글번호": m.group(1), "발행일": parsedate_to_datetime(pub)})
    return out


def compute_blog_metrics(rows, blog_cfg):
    """블로그 ID별 노출력 지표. 입력값은 전부 실측:
    순위(블로그 탭 실검색) + 발행 이력(공식 RSS) + 신규 글 진입(둘의 교집합)"""
    from datetime import datetime, timedelta, timezone
    now = datetime.now(timezone.utc)
    metrics = []
    for teacher, ids in blog_cfg["선생님"].items():
        for bid in ids:
            bid_l = bid.lower()
            hits = [d for r in rows for d in r.get("선생님상세", []) if d["블로그"] == bid_l]
            kw_best = {}
            for d in hits:
                kw_best[d["키워드"]] = min(kw_best.get(d["키워드"], 999), d["순위"])
            top10 = sum(1 for v in kw_best.values() if v <= 10)
            top30 = len(kw_best)
            avg_best = round(sum(kw_best.values()) / top30, 1) if top30 else None
            rss_err = ""
            try:
                posts = fetch_rss(bid)
            except Exception as e:
                posts, rss_err = [], f"RSS 실패({type(e).__name__})"
            pub30 = sum(1 for p in posts if (now - p["발행일"]).days <= 30)
            recent14 = [p for p in posts if (now - p["발행일"]).days <= 14]
            hit_posts = {d["글"] for d in hits}
            entered = sum(1 for p in recent14 if p["글번호"] in hit_posts)
            score = round(top10 * 4 + (top30 - top10) * 1.5 + pub30 * 1 + entered * 6, 1)
            metrics.append({"선생님": teacher, "블로그": bid, "30위내": top30, "10위내": top10,
                            "평균최고순위": avg_best, "발행30일": pub30, "새글14일": len(recent14),
                            "새글진입": entered, "점수": score, "비고": rss_err})
            time.sleep(0.2)
    return metrics


def write_teacher_sheet(metrics):
    """「선생님 지표」 시트 재작성 + 지표 기록들로 점수 추이 매트릭스"""
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    hist = {}  # 날짜 -> {블로그: 점수}
    for f in sorted(HISTORY_DIR.glob("지표_*.json")):
        with open(f, encoding="utf-8") as fp:
            hist[f.stem.replace("지표_", "")] = {m["블로그"]: m["점수"] for m in json.load(fp)}
    dates = sorted(hist)[-8:]
    prev = hist[dates[-2]] if len(dates) >= 2 else {}

    wb = load_workbook(XLSX_FILE)
    if "선생님 지표" in wb.sheetnames:
        del wb["선생님 지표"]
    ws = wb.create_sheet("선생님 지표", 1)

    pink = PatternFill("solid", fgColor="F8CBAD")
    thin = Border(*[Side(style="thin")] * 4)
    center = Alignment(horizontal="center", vertical="center")

    ws.cell(1, 1, "선생님 블로그 노출력 지표 — 입력값 전부 실측(블로그 탭 실제 순위 + 공식 RSS 발행 이력)").font = Font(bold=True)
    ws.cell(2, 1, "점수 = 10위내 키워드×4 + 11~30위 키워드×1.5 + 최근30일 발행×1 + 새글(14일내)의 30위 진입×6 · 절대평가 아닌 블로그 간 비교/주간 추이용")
    ws.cell(3, 1, "새글진입이 지수의 가장 민감한 신호: 지수 높은 블로그는 새 글이 바로 순위에 꽂히고, 지수가 내려가면 새 글부터 안 뜹니다")

    headers = ["선생님", "블로그", "커버(30위내)", "10위내", "평균 최고순위",
               "발행(30일)", "새글(14일)", "새글 진입", "점수", "전주 점수", "변동", "비고"]
    for c, h in enumerate(headers, 1):
        cell = ws.cell(5, c, h)
        cell.fill = pink
        cell.font = Font(bold=True)
        cell.alignment = center
        cell.border = thin

    for r, m in enumerate(sorted(metrics, key=lambda x: -x["점수"]), 6):
        old = prev.get(m["블로그"])
        change = "" if old is None else (f"▲{round(m['점수']-old,1)}" if m["점수"] > old
                                         else (f"▼{round(old-m['점수'],1)}" if m["점수"] < old else "="))
        vals = [m["선생님"], m["블로그"], m["30위내"], m["10위내"],
                m["평균최고순위"] if m["평균최고순위"] is not None else "",
                m["발행30일"], m["새글14일"], m["새글진입"], m["점수"],
                old if old is not None else "", change, m["비고"]]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(r, c, v)
            cell.border = thin
            cell.alignment = center

    # 점수 추이 매트릭스
    base = 6 + len(metrics) + 2
    ws.cell(base, 1, "■ 점수 추이 (최근 기록)").font = Font(bold=True)
    ws.cell(base + 1, 1, "블로그").fill = pink
    for c, d in enumerate(dates, 2):
        ws.cell(base + 1, c, d).fill = pink
    for r, m in enumerate(sorted(metrics, key=lambda x: -x["점수"]), base + 2):
        ws.cell(r, 1, m["블로그"])
        for c, d in enumerate(dates, 2):
            ws.cell(r, c, hist[d].get(m["블로그"], ""))

    for col, w in zip("ABCDEFGHIJKL", [10, 16, 12, 9, 13, 10, 10, 9, 8, 9, 8, 16]):
        ws.column_dimensions[col].width = w
    reorder_sheets(wb)
    wb.save(XLSX_FILE)


# ──────────────────────────── 4) 엑셀 출력 ────────────────────────────

def blog_cfg_cached():
    if not hasattr(blog_cfg_cached, "v"):
        with open(BLOGS_FILE, encoding="utf-8") as f:
            blog_cfg_cached.v = json.load(f)
    return blog_cfg_cached.v


def rank_str(ranks):
    return ",".join(str(r) for r in ranks) if ranks else ""


def load_previous():
    if not HISTORY_DIR.exists():
        return None, None
    files = sorted(f for f in HISTORY_DIR.glob("*.json")
                   if re.fullmatch(r"\d{4}-\d{2}-\d{2}", f.stem))
    files = [f for f in files if f.stem != date.today().isoformat()]
    if not files:
        return None, None
    with open(files[-1], encoding="utf-8") as f:
        return files[-1].stem, json.load(f)


def diff_mark(cur, prev):
    """순위 변동 표시: 숫자 작을수록 좋음. prev/cur 는 int 또는 None."""
    if cur is None:
        return f"이탈(전{prev})" if prev else ""
    if prev is None:
        return "신규" if cur else ""
    if cur < prev:
        return f"▲{prev - cur}"
    if cur > prev:
        return f"▼{cur - prev}"
    return ""


def fmt_ranks(ranks, mark):
    """순위 목록을 잘리지 않게 축약: 최대 4개 + 「외N」 + 변동 표시"""
    txt = ",".join(str(r) for r in ranks[:4])
    if len(ranks) > 4:
        txt += f" 외{len(ranks) - 4}"
    return (txt + (" " + mark if mark else "")).strip()


def reorder_sheets(wb):
    """선생님 지표 맨 앞 고정, 날짜 시트는 최신순"""
    pat = re.compile(r"(\d\d\.\d\d\.\d\d)(?: \((\d+)\))?$")
    teacher = [ws for ws in wb.worksheets if ws.title == "선생님 지표"]
    dated = [ws for ws in wb.worksheets if pat.match(ws.title)]
    dated.sort(key=lambda ws: (pat.match(ws.title).group(1),
                               int(pat.match(ws.title).group(2) or 1)), reverse=True)
    others = [ws for ws in wb.worksheets if ws not in teacher and ws not in dated]
    wb._sheets = teacher + dated + others
    if dated:
        wb.active = wb.worksheets.index(dated[0])  # 열면 최신 날짜 시트


def write_excel(rows, prev_date, prev_rows):
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    trainers = list(rows[0]["선생님_전체"]) if rows else []
    prev_map = {r["키워드"]: r for r in (prev_rows or [])}
    first_run = not prev_map  # 첫 실행(비교 기준 없음)이면 변동 표시 생략

    wb = load_workbook(XLSX_FILE) if XLSX_FILE.exists() else Workbook()
    if "Sheet" in wb.sheetnames and len(wb.sheetnames) == 1 and wb["Sheet"].max_row == 1:
        del wb["Sheet"]
    # 같은 날 다시 돌리면 덮어쓰지 않고 (2), (3) 시트로 추가 — 발행 후 재확인용
    base_name = date.today().strftime("%y.%m.%d")
    sheet_name, k = base_name, 2
    while sheet_name in wb.sheetnames:
        sheet_name = f"{base_name} ({k})"
        k += 1
    ws = wb.create_sheet(sheet_name, 0)

    yellow = PatternFill("solid", fgColor="FFFF00")
    pink = PatternFill("solid", fgColor="F8CBAD")
    green = PatternFill("solid", fgColor="00B050")
    red_font = Font(color="C00000", bold=True)
    thin = Border(*[Side(style="thin")] * 4)
    center = Alignment(horizontal="center", vertical="center")

    headers = (["no", "키워드", "검색량PC", "검색량MO", "검색량합계"]
               + trainers + ["리뷰어", "플레이스", "노출점수"])

    ws.cell(1, 1, "완성 시각").fill = yellow
    ws.cell(1, 2, datetime.now().strftime("%Y.%m.%d %H:%M")).fill = yellow
    ws.cell(1, 4, f"{센터표시()} (블로그 30위·플레이스 70위까지 추적, 플레이스는 광고 제외)").fill = yellow
    if prev_date:
        ws.cell(1, 3, f"비교: {prev_date}").fill = yellow

    for c, h in enumerate(headers, 1):
        cell = ws.cell(2, c, h)
        cell.fill = green if h in trainers else pink
        cell.font = Font(bold=True)
        cell.alignment = center
        cell.border = thin

    for r, row in enumerate(rows, 3):
        prev = prev_map.get(row["키워드"], {})
        prev_teachers = prev.get("선생님", {})
        best = lambda ranks: min(ranks) if ranks else None

        blue_f = Font(color="0070C0")
        down_f = Font(color="C00000")
        gray_f = Font(color="808080")

        blog_fail = any("블로그탭" in e for e in row["오류"])   # 이번 수집에서 블로그탭 캡차·오류
        vals = [(r - 2, ""), (row["키워드"], ""), (row["PC"], ""), (row["모바일"], ""), (row["합계"], "")]
        for t in trainers:
            cur_ranks = row["선생님"].get(t, [])
            # 블로그탭 측정 실패로 순위가 비면 '이탈'로 단정하지 않는다(측정 실패 ≠ 순위 하락)
            if first_run or (blog_fail and not cur_ranks):
                mark = ""
            else:
                mark = diff_mark(best(cur_ranks), best(prev_teachers.get(t, [])))
            vals.append((fmt_ranks(cur_ranks, mark), mark))
        vals.append((rank_str(row["리뷰어"][:4]) + (f" 외{len(row['리뷰어'])-4}" if len(row["리뷰어"]) > 4 else ""), ""))
        # 플레이스 변동 표시 — 측정 실패를 순위 변동으로 둔갑시키지 않는다
        cur_place_fail = any("플레이스" in e for e in row["오류"])   # 이번 수집에서 캡차·지도 오류
        prev_place = prev.get("플레이스")
        prev_place_fail = prev.get("플레이스오류", False)             # 지난주가 측정 실패였나(옛 기록엔 없을 수 있음)
        if first_run or (cur_place_fail and row["플레이스"] is None):
            p_mark = ""                              # 비교 기준 없음 / 이번 측정 실패 → 변동 판정 보류
        elif prev_place is None and row["플레이스"] is not None:
            # 지난주 값이 비어 있음: 측정 실패였는지 진짜 권외였는지 모르면 '신규' 단정 금지(순위만 표기)
            p_mark = "" if (prev_place_fail or "플레이스오류" not in prev) else diff_mark(row["플레이스"], prev_place)
        else:
            p_mark = diff_mark(row["플레이스"], prev_place)
        p_txt = "확인못함" if (cur_place_fail and row["플레이스"] is None) else (str(row["플레이스"]) if row["플레이스"] else "")
        vals.append(((p_txt + " " + p_mark).strip(), p_mark))
        vals.append((row["노출점수"], ""))

        for c, (v, mark) in enumerate(vals, 1):
            cell = ws.cell(r, c, v)
            cell.border = thin
            cell.alignment = center
            if "▲" in mark or mark == "신규":
                cell.font = blue_f
            elif "▼" in mark:
                cell.font = down_f
            elif "이탈" in mark:
                cell.font = gray_f

        # 구멍 키워드: 검색량 300+ 인데 선생님 글이 30위 안에 하나도 없음
        if row["합계"] >= 300 and not row["선생님"]:
            ws.cell(r, 2).font = red_font

        if row["오류"]:
            ws.cell(r, len(headers) + 1, " / ".join(row["오류"])).font = Font(color="808080", size=9)

    widths = [5, 18, 10, 10, 11] + [14] * len(trainers) + [15, 13, 10]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(2, i).column_letter].width = w
    ws.freeze_panes = "C3"

    # ── 표 옆 「이번 주 브리핑」 ──
    bc = len(headers) + 3  # 오류 칸 다음, 한 칸 띄우고
    ws.column_dimensions[ws.cell(2, bc).column_letter].width = 72
    section = PatternFill("solid", fgColor="DDEBF7")

    def put(text, bold=False, fill=None):
        put.r += 1
        cell = ws.cell(put.r, bc, text)
        cell.alignment = Alignment(wrap_text=False, vertical="center")
        if bold:
            cell.font = Font(bold=True)
        if fill:
            cell.fill = fill
    put.r = 1

    best_rank = lambda row: min((r for v in row["선생님"].values() for r in v), default=None)

    put("■ 이 표 읽는 법", True, section)
    put("· 노출점수 = 검색량 × 순위별 클릭률 추정. 이 키워드가 한 달에 데려오는 사람 수(키워드끼리 비교용)")
    put(f"· 리뷰어 = 선생님 외 글 중 {센터표시()} 언급·플레이스 지도 첨부가 확인된 글의 순위")
    put("· 플레이스 = 광고 제외 순위(70위까지) / 블로그는 30위까지 추적")
    put("· 빨간 키워드 = 검색량 300+ 인데 선생님 글이 30위 안에 없음 → 글감 후보")
    put("· ▲2=2계단 상승(파랑) · ▼3=하락(빨강) · 신규=새로 진입 · 이탈(전7)=지난번 7위였는데 순위권 밖(회색)")
    put("· 순위가 5개 이상이면 「1,2,3,4 외2」로 줄여 표시 · 같은 날 다시 돌리면 (2) 시트로 추가됩니다")
    put("· 「확인못함」 = 그때 캡차·지도 오류로 못 쟀다는 뜻(순위 하락 아님) · 변동(▲▼신규이탈)은 맨 위 「비교:날짜」 기준")
    put("")

    put("■ 이번 주 글감 후보 (기대 유입 큰 순 · 난이도는 검색량 기준 추정)", True, section)
    cands = sorted((r for r in rows
                    if r["합계"] >= 100 and (best_rank(r) is None or best_rank(r) > 10)),
                   key=lambda r: -r["합계"])
    for r in cands[:8]:
        gain = round(r["합계"] * ctr_of(5))
        diff = "상" if r["합계"] >= 1000 else ("중" if r["합계"] >= 300 else "하")
        notes = []
        rv10 = [x for x in r["리뷰어"] if x <= 10]
        if rv10:
            notes.append(f"체험단 글 {min(rv10)}위 = 일반 블로그도 뚫리는 판(진입 검증)")
        br = best_rank(r)
        if br and 11 <= br <= 30:
            who = next(name for name, v in r["선생님"].items() if br in v)
            notes.append(f"{who} {br}위 글 있음 → 보강·재발행 or 새 글")
        put(f"· {r['키워드']} — 월 {r['합계']:,}회 | 5위 진입 시 +{gain}명/월 | 난이도 {diff}"
            + (" | " + " / ".join(notes) if notes else ""))
    if not cands:
        put("· 없음 — 검색량 있는 키워드에 전부 10위 안 글이 있습니다")
    put("")

    put("■ 신규 키워드 후보 (검색광고 연관 키워드에서 자동 추출)", True, section)
    rel_file = HISTORY_DIR / "_연관키워드.json"
    have = {norm(r["키워드"]) for r in rows}
    region_words = blog_cfg_cached().get("지역_단어", [])
    biz_words = [w.lower() for w in blog_cfg_cached().get("업종_단어", [])]
    exclude_words = [w.lower() for w in blog_cfg_cached().get("제외_단어", [])]
    own_name = blog_cfg_cached().get("플레이스_센터명", "")

    # 수집 때 본 업체명들 — 타 센터 상호 자동 판별용
    shop_file = HISTORY_DIR / "_플레이스상호.json"
    shop_norm = []
    if shop_file.exists():
        with open(shop_file, encoding="utf-8") as f:
            shop_norm = [re.sub(r"[\s&/·|()\-]", "", n).lower()
                         for n in json.load(f) if own_name not in n]
    strip_tokens = sorted(region_words + [w for w in biz_words], key=len, reverse=True)

    def is_brand(kwtext):
        """타 센터 상호/금지어 판별: 지역·업종 단어를 걷어낸 알맹이가 실제 업체명에 들어 있으면 상호"""
        t = kwtext.replace(" ", "").lower()
        if any(w in t for w in exclude_words):
            return True
        core = t
        for tok in strip_tokens:
            core = core.replace(tok.lower(), "")
        core = core.strip()
        return len(core) >= 2 and any(core in s for s in shop_norm)

    news, brand_cut = [], 0
    if rel_file.exists():
        with open(rel_file, encoding="utf-8") as f:
            pool = json.load(f)
        for x in pool:
            if (x["합계"] >= 50 and norm(x["키워드"]) not in have
                    and any(w in x["키워드"].replace(" ", "") for w in region_words)
                    and any(w in x["키워드"].replace(" ", "").lower() for w in biz_words)):
                if is_brand(x["키워드"]):
                    brand_cut += 1
                else:
                    news.append(x)
    if news:
        for x in news[:10]:
            comp = {"높음": "광고경쟁 높음", "중간": "광고경쟁 중간", "낮음": "광고경쟁 낮음"}.get(x["광고경쟁"], "")
            put(f"· {x['키워드']} — 월 {x['합계']:,}회 (PC {x['PC']} + 모바일 {x['모바일']})"
                + (f" | {comp}" if comp else ""))
        put("  → 추적에 넣으려면 설정_키워드.txt 에 한 줄 추가하면 다음 실행부터 잡힙니다")
        if brand_cut:
            put(f"  → 타 센터 상호·금지어로 보이는 후보 {brand_cut}개는 자동 제외 (설정_블로그.json 제외_단어에서 조정)")
    else:
        put("· 지역 연관 키워드 중 검색량 50+ 신규 후보 없음")
    put("")

    put("■ 체험단 의존 주의 (상위가 외부 글뿐 — 내려가면 순위 통째로 빠짐)", True, section)
    risky = [r for r in rows
             if r["리뷰어"] and min(r["리뷰어"]) <= 10
             and (best_rank(r) is None or best_rank(r) > min(r["리뷰어"]))]
    if risky:
        for r in sorted(risky, key=lambda x: -x["합계"])[:5]:
            br = best_rank(r)
            put(f"· {r['키워드']} — 외부 글 {min(r['리뷰어'])}위가 최상위"
                + (f", 선생님 글은 {br}위" if br else ", 선생님 글 없음")
                + " → 선생님 글로 백업 필요")
    else:
        put("· 해당 없음")
    put("")

    put("■ 지난 기록 대비 변동", True, section)
    if first_run:
        put("· 첫 기록입니다 — 다음 실행부터 급등·급락이 여기 표시됩니다")
    else:
        moves = []
        for r in rows:
            prev = prev_map.get(r["키워드"], {})
            for cur, old, label in [
                (best_rank(r), min((x for v in prev.get("선생님", {}).values() for x in v), default=None), "블로그"),
                (r["플레이스"], prev.get("플레이스"), "플레이스"),
            ]:
                if cur and old and abs(cur - old) >= 3:
                    arrow = "▲" if cur < old else "▼"
                    moves.append((abs(cur - old), f"· {r['키워드']} {label} {old}위 → {cur}위 {arrow}{abs(cur-old)}"))
                elif old and not cur:  # 순위권에서 사라짐 — 조용히 넘기지 않는다
                    limit = "70위" if label == "플레이스" else "30위"
                    moves.append((100, f"· {r['키워드']} {label} {old}위 → {limit} 밖 이탈 ⚠"))
        if moves:
            for _, line in sorted(moves, reverse=True)[:8]:
                put(line)
        else:
            put("· 3계단 이상 움직인 순위 없음")

    reorder_sheets(wb)
    wb.save(XLSX_FILE)


# ──────────────────────────── main ────────────────────────────

def main():
    with open(API_CONFIG, encoding="utf-8") as f:
        api_cfg = json.load(f)
    with open(BLOGS_FILE, encoding="utf-8") as f:
        blog_cfg = json.load(f)
    keywords = list(dict.fromkeys(  # 중복 줄 제거 (순서 유지)
        ln.strip() for ln in KEYWORDS_FILE.read_text(encoding="utf-8").splitlines() if ln.strip()))
    test_mode = len(sys.argv) > 1
    if test_mode:
        keywords = [k for k in keywords if k in sys.argv[1:]] or sys.argv[1:]

    # 실행 흔적 남기기 — 중간에 창이 닫혀도 시작 여부는 진단 가능
    HISTORY_DIR.mkdir(exist_ok=True)
    with open(HISTORY_DIR / "_마지막실행.txt", "a", encoding="utf-8") as f:
        import datetime
        f.write(f"{datetime.datetime.now():%Y-%m-%d %H:%M:%S} 시작 (키워드 {len(keywords)}개)\n")

    print(f"■ 키워드 {len(keywords)}개")
    print("■ 1/3 검색량 조회 …")
    try:
        volumes, related = fetch_volumes(keywords, api_cfg)
        print(f"  검색량 확보: {len(volumes)}/{len(keywords)} (연관 키워드 풀 {len(related)}개)")
        if not test_mode:
            HISTORY_DIR.mkdir(exist_ok=True)
            with open(HISTORY_DIR / "_연관키워드.json", "w", encoding="utf-8") as f:
                json.dump(sorted(related.values(), key=lambda x: -x["합계"]),
                          f, ensure_ascii=False, indent=0)
    except Exception as e:
        volumes = {}
        print(f"  ⚠ 검색량 조회 실패({type(e).__name__}) — 순위 수집은 계속하고 검색량만 0으로 둡니다")
        print("    (인터넷 연결과 _네이버광고_API.json 의 키를 확인하세요)")

    print("■ 2/3 블로그 탭 + 플레이스 수집 …")
    from playwright.sync_api import sync_playwright
    with sync_playwright() as p:
        browser = p.chromium.launch(headless=True)
        page = browser.new_page(viewport={"width": 1400, "height": 900})
        serp = collect_serp(keywords, page)
        browser.close()

    if not test_mode:
        # 이번 수집에서 본 업체명 전부 저장 — 신규 키워드 후보에서 타 센터 상호 거르는 재료
        names = sorted({pl["name"] for e in serp.values() for pl in e["places"]})
        with open(HISTORY_DIR / "_플레이스상호.json", "w", encoding="utf-8") as f:
            json.dump(names, f, ensure_ascii=False, indent=0)

    print("■ 3/3 판정 및 저장 … (제목에 센터명 없는 글은 본문을 직접 확인)")
    HISTORY_DIR.mkdir(exist_ok=True)
    cache = {}
    if CACHE_FILE.exists():
        with open(CACHE_FILE, encoding="utf-8") as f:
            cache = json.load(f)
    stats = {"조회": 0, "실패": 0}
    trainer_names = list(blog_cfg["선생님"].keys())
    rows = []
    for i, kw in enumerate(keywords, 1):
        row = analyze(kw, serp[kw], volumes, blog_cfg, cache, stats)
        row["선생님_전체"] = trainer_names
        rows.append(row)
        print(f"  ({i}/{len(keywords)}) {kw} 판정 완료 (본문조회 누적 {stats['조회']}건)")
    with open(CACHE_FILE, "w", encoding="utf-8") as f:
        json.dump(cache, f, ensure_ascii=False, indent=0)
    print(f"  본문 확인: 새로 {stats['조회']}건 (캐시 {len(cache)}건, 실패 {stats['실패']}건)")
    hist_file = HISTORY_DIR / ("_테스트.json" if test_mode else f"{date.today().isoformat()}.json")
    with open(hist_file, "w", encoding="utf-8") as f:
        json.dump(rows, f, ensure_ascii=False, indent=1)

    if test_mode:
        print(f"\n[테스트 모드] 엑셀은 건드리지 않음 → {hist_file.name} 만 저장")
    else:
        prev_date, prev_rows = load_previous()
        try:
            write_excel(rows, prev_date, prev_rows)
            print("■ 선생님 블로그 지표 계산 … (RSS 발행 이력 수집)")
            metrics = compute_blog_metrics(rows, blog_cfg)
            with open(HISTORY_DIR / f"지표_{date.today().isoformat()}.json", "w", encoding="utf-8") as f:
                json.dump(metrics, f, ensure_ascii=False, indent=1)
            write_teacher_sheet(metrics)
        except PermissionError:
            print("\n⚠ 키워드순위표.xlsx 가 열려 있어서 저장하지 못했습니다.")
            print("  엑셀을 닫은 뒤 실행.bat 을 다시 눌러주세요. (수집한 순위는 기록 폴더에 남아 있습니다)")
            return
        print(f"\n완료 → {XLSX_FILE.name} (시트 {date.today().strftime('%y.%m.%d')} + 선생님 지표)")
    holes = [r for r in rows if r["합계"] >= 300 and not r["선생님"]]
    if holes:
        print("\n※ 구멍 키워드 (검색량 300+ 인데 선생님 글 30위 안에 없음):")
        for r in sorted(holes, key=lambda x: -x["합계"]):
            print(f"   {r['키워드']}: 월 {r['합계']:,}회")
    errs = [(r["키워드"], r["오류"]) for r in rows if r["오류"]]
    if errs:
        print("\n※ 수집 오류:")
        for kw, e in errs:
            print(f"   {kw}: {e}")


if __name__ == "__main__":
    main()
